"""Render validated resume content into a one-page A4 .docx (python-docx)."""
import io
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: F401  (kept for users extending the layout)
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x38, 0x64)
FONT = "Calibri"


def slug(s: str, n: int = 40) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", s or "").strip("_")[:n]


def _run(p, text, size=10, bold=False, italic=False, color=None):
    r = p.add_run(text)
    r.font.name, r.font.size, r.bold, r.italic = FONT, Pt(size), bold, italic
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if color:
        r.font.color.rgb = color
    return r


def _para(doc, before=0, after=2, align=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), 1.0
    if align is not None:
        p.alignment = align
    return p


def _rule(p):
    """Bottom border under a section heading."""
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in {"w:val": "single", "w:sz": "6", "w:space": "1", "w:color": "1F3864"}.items():
        bottom.set(qn(k), v)
    bdr.append(bottom)
    pPr.append(bdr)


def _section(doc, title):
    p = _para(doc, before=5, after=2)
    _run(p, title.upper(), size=10.5, bold=True, color=NAVY)
    _rule(p)


def _bullet(doc, text):
    p = _para(doc, after=1)
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.first_line_indent = Cm(-0.3)
    _run(p, "•  " + text, size=10)


def _right_tab(p, width_cm=18.6):
    p.paragraph_format.tab_stops.add_tab_stop(Cm(width_cm), WD_TAB_ALIGNMENT.RIGHT)


def build_docx(profile_raw: dict, content: dict) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.0)
    sec.left_margin = sec.right_margin = Cm(1.2)

    p = _para(doc, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    _run(p, profile_raw["name"].upper(), size=16, bold=True)
    p = _para(doc, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    _run(p, content["headline"], size=10.5, bold=True, color=NAVY)
    contact = [profile_raw["location"], profile_raw["email"], profile_raw["phone"], profile_raw.get("linkedin") or ""]
    p = _para(doc, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    _run(p, "  •  ".join(c for c in contact if c), size=9)

    _section(doc, "Professional Summary")
    p = _para(doc, after=1)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _run(p, content["summary"], size=10)

    _section(doc, "Core Skills")
    p = _para(doc, after=1)
    _run(p, "  •  ".join(content["skills"]), size=10)

    _section(doc, "Professional Experience")
    for role in profile_raw["experience"]:
        p = _para(doc, before=3, after=1)
        _right_tab(p)
        _run(p, role["title"], size=10, bold=True)
        _run(p, f"  |  {role['company']}, {role['location']}", size=10)
        _run(p, f"\t{role['dates']}", size=9.5, italic=True)
        for text in content["bullets_by_role"].get(role["id"], []):
            _bullet(doc, text)

    _section(doc, "Certifications")
    for c in profile_raw.get("certifications", []):
        _bullet(doc, c)

    _section(doc, "Education")
    for e in profile_raw.get("education", []):
        p = _para(doc, after=1)
        _right_tab(p)
        _run(p, f"{e['degree']}  |  {e['school']}", size=9.5, bold=True)
        _run(p, f"\t{e['year']}", size=10)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def file_name(profile_raw: dict, company: str, title: str) -> str:
    return f"{slug(profile_raw['name'], 30)}_{slug(company, 25)}_{slug(title, 35)}.docx"
